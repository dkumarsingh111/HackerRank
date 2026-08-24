"""
HPPCS01 - AI Resume Tailoring & ATS Optimization
================================================

Inputs
------
1. Resume: PDF or Word (.docx / .doc)
2. Job description: pasted text or a text file

Local LLMs through Ollama
-------------------------
- Gemma 3 4B  : extraction + JD parsing + tailoring + revision
- Llama 3.2 1B: independent review

Outputs
-------
- resume_source_text.txt
- resume_profile.json
- job_description.json
- ats_analysis.json
- tailored_resume.txt
- tailored_resume.docx
- tailored_resume.pdf   (when LibreOffice is installed)
- final_review.json

Design principles
-----------------
- The input resume can be unstructured.
- The LLM may improve wording and structure.
- The LLM must NOT invent candidate facts.
- Achievements/projects are included only when present in the source.
- Missing JD keywords are reported, not fabricated into the resume.
- Python validation is used in addition to the independent Llama review.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Tuple

import requests

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:
    Document = None
    WD_ALIGN_PARAGRAPH = None
    Inches = Pt = None


DEFAULT_GEMMA_MODEL = "gemma3:4b"
DEFAULT_REVIEW_MODEL = "llama3.2:1b"
DEFAULT_OLLAMA_HOST = "http://localhost:11434"

SECTION_NAMES = {
    "PROFESSIONAL SUMMARY",
    "SUMMARY",
    "PROFILE",
    "SKILLS",
    "PROFESSIONAL EXPERIENCE",
    "EXPERIENCE",
    "WORK EXPERIENCE",
    "EDUCATION",
    "ACHIEVEMENTS",
    "PROJECTS",
    "INTERESTS",
    "CERTIFICATIONS",
}


# ---------------------------------------------------------------------------
# Document parsing
# ---------------------------------------------------------------------------

def extract_pdf_text(path: Path) -> str:
    if pdfplumber is None:
        raise RuntimeError(
            "pdfplumber is required for PDF input. "
            "Install it with: pip install pdfplumber"
        )

    pages: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {page_number}]\n{text}")

    result = "\n\n".join(pages).strip()

    if not result:
        raise RuntimeError(
            f"No text could be extracted from {path.name}. "
            "The PDF may be scanned/image-only. OCR support is required "
            "for image-only PDFs."
        )

    return result


def extract_docx_text(path: Path) -> str:
    if Document is None:
        raise RuntimeError(
            "python-docx is required for Word input. "
            "Install it with: pip install python-docx"
        )

    doc = Document(path)
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"[Table {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    result = "\n".join(parts).strip()

    if not result:
        raise RuntimeError(f"No text could be extracted from {path.name}.")

    return result


def extract_doc_text(path: Path) -> str:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "Legacy .doc input requires LibreOffice/soffice in PATH."
        )

    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                tmp,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )

        if result.returncode != 0:
            raise RuntimeError(
                result.stderr.strip() or "Could not convert DOC to text."
            )

        candidates = list(Path(tmp).glob("*.txt"))
        if not candidates:
            raise RuntimeError("LibreOffice produced no TXT output.")

        return candidates[0].read_text(
            encoding="utf-8", errors="replace"
        ).strip()


def extract_resume_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    if suffix == ".doc":
        return extract_doc_text(path)

    raise ValueError(
        f"Unsupported resume format: {suffix}. "
        "Use PDF, DOCX or DOC."
    )


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clean_list(value: Any) -> List[str]:
    if value is None:
        return []

    if isinstance(value, str):
        value = [value]

    if not isinstance(value, list):
        return []

    result: List[str] = []

    for item in value:
        if isinstance(item, dict):
            item = " | ".join(
                f"{k}: {v}"
                for k, v in item.items()
                if v not in (None, "")
            )

        item = clean_text(item)

        if item and item.casefold() not in {
            "none",
            "n/a",
            "not provided",
            "unknown",
        }:
            result.append(item)

    seen = set()
    unique = []

    for item in result:
        key = item.casefold()
        if key not in seen:
            seen.add(key)
            unique.append(item)

    return unique


def extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()

    text = re.sub(r"^```json\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)

    try:
        obj = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.S)

        if not match:
            raise ValueError("LLM did not return valid JSON.")

        obj = json.loads(match.group(0))

    if not isinstance(obj, dict):
        raise ValueError("Expected a JSON object from LLM.")

    return obj


# ---------------------------------------------------------------------------
# Ollama
# ---------------------------------------------------------------------------

def call_ollama(
    model: str,
    prompt: str,
    host: str,
    timeout: int,
    num_predict: int,
    json_mode: bool = False,
) -> str:

    payload: Dict[str, Any] = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "keep_alive": "5m",
        "options": {
            "temperature": 0.1,
            "top_p": 0.9,
            "num_predict": num_predict,
        },
    }

    if json_mode:
        payload["format"] = "json"

    response = requests.post(
        f"{host.rstrip('/')}/api/generate",
        json=payload,
        timeout=timeout,
    )

    response.raise_for_status()

    result = response.json().get("response", "")

    if not result.strip():
        raise RuntimeError(f"{model} returned an empty response.")

    return result.strip()


# ---------------------------------------------------------------------------
# Deterministic source anchors
# ---------------------------------------------------------------------------

def extract_source_anchors(text: str) -> Dict[str, str]:
    """Protect obvious contact facts from small-model extraction errors."""

    email_match = re.search(
        r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b",
        text,
    )

    email = email_match.group(0) if email_match else ""

    phone = ""

    phone_patterns = [
        r"(?<!\d)(\+\d{1,3}[\s-]?\d{3,5}[\s-]?\d{4,6})(?!\d)",
        r"(?<!\d)(\d{10})(?!\d)",
        r"(?<!\d)(\d{3}[\s-]\d{4})(?!\d)",
    ]

    for pattern in phone_patterns:
        match = re.search(pattern, text)
        if match:
            phone = match.group(1)
            break

    lines = [clean_text(x) for x in text.splitlines() if clean_text(x)]

    name = ""

    for line in lines[:8]:
        match = re.match(
            r"^(?:name|candidate)\s*[:\-]\s*(.+)$",
            line,
            flags=re.I,
        )

        if match:
            name = clean_text(match.group(1))
            break

    if not name:
        for line in lines[:5]:
            if (
                1 <= len(line.split()) <= 5
                and "@" not in line
                and not re.search(r"\d", line)
                and not re.search(
                    r"\b(summary|profile|skills|experience|education)\b",
                    line,
                    flags=re.I,
                )
            ):
                name = line
                break

    return {
        "name": name,
        "email": email,
        "phone": phone,
    }


# ---------------------------------------------------------------------------
# Resume extraction
# ---------------------------------------------------------------------------

def normalize_resume_profile(
    data: Dict[str, Any],
    anchors: Dict[str, str],
) -> Dict[str, Any]:

    profile: Dict[str, Any] = {
        "name": clean_text(data.get("name")),
        "contact": {
            "email": clean_text(
                (data.get("contact") or {}).get("email")
            ),
            "phone": clean_text(
                (data.get("contact") or {}).get("phone")
            ),
            "location": clean_text(
                (data.get("contact") or {}).get("location")
            ),
        },
        "education": clean_list(data.get("education")),
        "experience": clean_list(data.get("experience")),
        "skills": clean_list(data.get("skills")),
        "achievements": clean_list(data.get("achievements")),
        "projects": clean_list(data.get("projects")),
        "interests": clean_list(data.get("interests")),
    }

    # Source anchors override questionable LLM values.
    if anchors.get("name"):
        profile["name"] = anchors["name"]

    if anchors.get("email"):
        profile["contact"]["email"] = anchors["email"]

    if anchors.get("phone"):
        profile["contact"]["phone"] = anchors["phone"]

    return profile


def extract_resume_profile(
    resume_text: str,
    model: str,
    host: str,
    timeout: int,
) -> Dict[str, Any]:

    anchors = extract_source_anchors(resume_text)

    prompt = f"""
You are a resume FACT EXTRACTION engine.

The resume below can be completely unstructured:
paragraphs, bullets, tables, labels, sentences, or mixed formatting.

Extract ONLY facts explicitly supported by the SOURCE.

Return JSON only with exactly these keys:

{{
  "name": "",
  "contact": {{
    "email": "",
    "phone": "",
    "location": ""
  }},
  "education": [],
  "experience": [],
  "skills": [],
  "achievements": [],
  "projects": [],
  "interests": []
}}

Rules:
- Never invent information.
- Do not infer a job title from an employer name.
- Do not infer responsibilities from a job title.
- Preserve employers and dates.
- Preserve explicit achievements and metrics.
- Preserve explicit project names/details.
- Do not move responsibilities into achievements.
- If achievements are not present, return [].
- If projects are not present, return [].
- If interests are not present, return [].
- Do not turn years such as 2020-2022 into phone numbers.
- Do not infer location.
- Do not add certifications or technologies not in the source.

SOURCE:
{resume_text}
"""

    try:
        raw_result = call_ollama(
            model=model,
            prompt=prompt,
            host=host,
            timeout=timeout,
            num_predict=650,
            json_mode=True,
        )

        data = extract_json(raw_result)

    except Exception as exc:
        print(f"[WARN] Resume extraction failed: {exc}")
        data = {}

    return normalize_resume_profile(data, anchors)


# ---------------------------------------------------------------------------
# Job description extraction
# ---------------------------------------------------------------------------

def parse_job_description(
    jd_text: str,
    model: str,
    host: str,
    timeout: int,
) -> Dict[str, Any]:

    prompt = f"""
You are a JOB DESCRIPTION extraction engine.

Extract the job posting into structured JSON.

Return exactly:

{{
  "job_title": "",
  "company": "",
  "requirements": [],
  "responsibilities": [],
  "preferred_qualifications": [],
  "keywords": []
}}

Rules:
- Extract only information present in the job description.
- Do not invent requirements.
- Keywords should include explicit technical skills, tools,
  certifications, job titles, domain terms and important phrases.
- Remove duplicate keywords.
- Keep requirements and responsibilities separate.

JOB DESCRIPTION:
{jd_text}
"""

    result = call_ollama(
        model=model,
        prompt=prompt,
        host=host,
        timeout=timeout,
        num_predict=500,
        json_mode=True,
    )

    data = extract_json(result)

    return {
        "job_title": clean_text(data.get("job_title")),
        "company": clean_text(data.get("company")),
        "requirements": clean_list(data.get("requirements")),
        "responsibilities": clean_list(data.get("responsibilities")),
        "preferred_qualifications": clean_list(
            data.get("preferred_qualifications")
        ),
        "keywords": clean_list(data.get("keywords")),
    }


# ---------------------------------------------------------------------------
# ATS matching
# ---------------------------------------------------------------------------

def build_resume_search_text(profile: Dict[str, Any]) -> str:
    values: List[str] = [
        profile.get("name", ""),
        profile.get("contact", {}).get("location", ""),
    ]

    for field in (
        "education",
        "experience",
        "skills",
        "achievements",
        "projects",
        "interests",
    ):
        values.extend(profile.get(field, []))

    return " ".join(values).casefold()


def keyword_match(
    keywords: List[str],
    resume_text: str,
) -> Tuple[List[str], List[str]]:

    matched = []
    missing = []

    for keyword in keywords:
        key = clean_text(keyword)

        if not key:
            continue

        if key.casefold() in resume_text:
            matched.append(key)
        else:
            missing.append(key)

    return matched, missing


def calculate_ats_analysis(
    profile: Dict[str, Any],
    jd: Dict[str, Any],
) -> Dict[str, Any]:

    resume_text = build_resume_search_text(profile)

    keywords = clean_list(jd.get("keywords"))

    matched, missing = keyword_match(
        keywords,
        resume_text,
    )

    keyword_score = (
        round((len(matched) / len(keywords)) * 100)
        if keywords
        else 100
    )

    required_text = " ".join(jd.get("requirements", []))
    required_keywords = clean_list(
        re.findall(
            r"\b[A-Za-z][A-Za-z0-9+#./-]{1,30}\b",
            required_text,
        )
    )

    required_matched, required_missing = keyword_match(
        required_keywords,
        resume_text,
    )

    required_score = (
        round(
            len(required_matched)
            / len(required_keywords)
            * 100
        )
        if required_keywords
        else 100
    )

    return {
        "keyword_score": keyword_score,
        "required_skill_score": required_score,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "matched_requirements": required_matched,
        "missing_requirements": required_missing,
        "recommendations": [
            (
                f"Consider adding '{item}' only if you genuinely "
                "have this skill or experience."
            )
            for item in missing[:10]
        ],
    }


# ---------------------------------------------------------------------------
# Tailored CV generation
# ---------------------------------------------------------------------------

def generate_tailored_resume(
    profile: Dict[str, Any],
    jd: Dict[str, Any],
    ats: Dict[str, Any],
    model: str,
    host: str,
    timeout: int,
    revision_instruction: str = "",
) -> str:

    prompt = f"""
You are a professional resume writer.

Create a polished, ATS-friendly resume using the CANDIDATE FACTS and
JOB DESCRIPTION below.

The candidate resume may originally have been completely unstructured.

Your task:
1. Reorganize information professionally.
2. Improve grammar and wording.
3. Align relevant experience with the target job.
4. Use relevant keywords ONLY when the candidate facts support them.
5. Do not fabricate missing skills.
6. Do not fabricate achievements.
7. Do not fabricate metrics.
8. Do not fabricate employers, titles or dates.
9. Do not fabricate projects.
10. Do not create sections with no source data.

CRITICAL:
A missing JD keyword must NOT be inserted as a candidate skill.
Only report it in the ATS analysis.

Required resume structure:

CANDIDATE NAME
Phone | Email | Location

PROFESSIONAL SUMMARY

SKILLS

PROFESSIONAL EXPERIENCE

EDUCATION

Optional only if source data exists:
ACHIEVEMENTS
PROJECTS
CERTIFICATIONS
INTERESTS

For experience:
- Use company, job title and dates when supplied.
- Convert source responsibilities into concise professional bullets.
- Do not add responsibilities that are not supported.

For the summary:
- Write 2-3 sentences.
- Base every claim on candidate facts.
- Do not claim a level of expertise that is not stated.

Return plain text only.
Do not return Markdown code fences.
Do not return explanations.

CANDIDATE FACTS:
{json.dumps(profile, indent=2, ensure_ascii=False)}

JOB DESCRIPTION:
{json.dumps(jd, indent=2, ensure_ascii=False)}

ATS ANALYSIS:
{json.dumps(ats, indent=2, ensure_ascii=False)}

USER REVISION REQUEST:
{revision_instruction or "None"}
"""

    return call_ollama(
        model=model,
        prompt=prompt,
        host=host,
        timeout=timeout,
        num_predict=1000,
        json_mode=False,
    ).strip()


# ---------------------------------------------------------------------------
# User review / revision
# ---------------------------------------------------------------------------

def revise_resume(
    current_resume: str,
    profile: Dict[str, Any],
    jd: Dict[str, Any],
    user_instruction: str,
    model: str,
    host: str,
    timeout: int,
) -> str:

    prompt = f"""
Revise the CURRENT RESUME according to the USER REQUEST.

Rules:
- Keep all candidate facts grounded in CANDIDATE FACTS.
- Never invent missing information.
- Never add a missing skill simply because it appears in the JD.
- Keep employer names, dates, education and metrics accurate.
- If the user requests a change that would require an unsupported fact,
  make the best truthful wording possible.
- Return the complete revised resume as plain text.

CANDIDATE FACTS:
{json.dumps(profile, indent=2, ensure_ascii=False)}

JOB DESCRIPTION:
{json.dumps(jd, indent=2, ensure_ascii=False)}

CURRENT RESUME:
{current_resume}

USER REQUEST:
{user_instruction}
"""

    return call_ollama(
        model=model,
        prompt=prompt,
        host=host,
        timeout=timeout,
        num_predict=1000,
        json_mode=False,
    ).strip()


# ---------------------------------------------------------------------------
# Deterministic validation
# ---------------------------------------------------------------------------

def validate_grounding(
    resume: str,
    profile: Dict[str, Any],
) -> Dict[str, Any]:

    lower_resume = resume.casefold()
    issues: List[str] = []

    contact = profile.get("contact", {})

    protected_values = [
        ("name", profile.get("name", "")),
        ("email", contact.get("email", "")),
        ("phone", contact.get("phone", "")),
    ]

    missing = []

    for label, value in protected_values:
        if value and value.casefold() not in lower_resume:
            missing.append(label)

    if missing:
        issues.append(
            "Missing source contact information: "
            + ", ".join(missing)
        )

    # Every explicit skill should normally survive the tailoring.
    for skill in profile.get("skills", []):
        if skill.casefold() not in lower_resume:
            issues.append(f"Source skill missing: {skill}")

    # Explicit project and achievement phrases are important.
    for project in profile.get("projects", []):
        # Compare significant words rather than requiring the whole sentence.
        words = [
            w.casefold()
            for w in re.findall(r"[A-Za-z0-9+#.-]+", project)
            if len(w) > 3
        ]

        if words:
            present = sum(w in lower_resume for w in words)
            if present / len(words) < 0.5:
                issues.append(f"Project may be missing: {project}")

    for achievement in profile.get("achievements", []):
        numbers = re.findall(r"\d+(?:\.\d+)?%?", achievement)

        for number in numbers:
            if number not in resume:
                issues.append(
                    f"Achievement metric missing: {number}"
                )

    # Sections with no source data should not appear.
    empty_section_rules = [
        ("ACHIEVEMENTS", "achievements"),
        ("PROJECTS", "projects"),
        ("INTERESTS", "interests"),
    ]

    for heading, field in empty_section_rules:
        if (
            not profile.get(field)
            and re.search(rf"\b{heading}\b", resume, re.I)
        ):
            issues.append(
                f"{heading} section exists although source data is empty."
            )

    placeholders = re.findall(
        r"\[[^\]]+\]|\b(?:N/A|TBD|TO BE PROVIDED)\b",
        resume,
        flags=re.I,
    )

    if placeholders:
        issues.append(
            "Placeholder text found: "
            + ", ".join(placeholders[:5])
        )

    score = max(
        0,
        100 - min(100, len(issues) * 10),
    )

    return {
        "score": score,
        "issues": issues,
        "pass": not issues,
    }


# ---------------------------------------------------------------------------
# Independent Llama reviewer
# ---------------------------------------------------------------------------

def review_resume(
    resume: str,
    source_text: str,
    profile: Dict[str, Any],
    jd: Dict[str, Any],
    ats: Dict[str, Any],
    model: str,
    host: str,
    timeout: int,
) -> Dict[str, Any]:

    prompt = f"""
You are an independent resume quality reviewer.

Review the generated resume against:
1. ORIGINAL SOURCE
2. STRUCTURED CANDIDATE FACTS
3. JOB DESCRIPTION

Return JSON only:

{{
  "overall_score": 0,
  "ats_score": 0,
  "factual_consistency": 0,
  "keyword_alignment": 0,
  "format_quality": 0,
  "missing_source_items": [],
  "unsupported_claims": [],
  "issues": [],
  "recommendation": "PASS"
}}

Use "REGENERATE" when there is a meaningful factual or structural problem.

Important:
- The source is the ultimate truth.
- Professional rewriting is allowed.
- Invented experience is not allowed.
- Invented skills are not allowed.
- Invented metrics are not allowed.
- Missing source achievements/projects are problems.
- Empty source achievements/projects are NOT problems.
- A missing JD keyword is NOT a hallucination; it should be reported as
  missing rather than fabricated.

ORIGINAL SOURCE:
{source_text}

STRUCTURED CANDIDATE FACTS:
{json.dumps(profile, indent=2, ensure_ascii=False)}

JOB DESCRIPTION:
{json.dumps(jd, indent=2, ensure_ascii=False)}

ATS ANALYSIS:
{json.dumps(ats, indent=2, ensure_ascii=False)}

GENERATED RESUME:
{resume}
"""

    try:
        result = extract_json(
            call_ollama(
                model=model,
                prompt=prompt,
                host=host,
                timeout=timeout,
                num_predict=550,
                json_mode=True,
            )
        )
    except Exception as exc:
        return {
            "overall_score": 0,
            "ats_score": ats.get("keyword_score", 0),
            "factual_consistency": 0,
            "keyword_alignment": ats.get("keyword_score", 0),
            "format_quality": 0,
            "missing_source_items": [],
            "unsupported_claims": [],
            "issues": [f"Llama review failed: {exc}"],
            "recommendation": "REGENERATE",
        }

    return result


# ---------------------------------------------------------------------------
# DOCX/PDF generation
# ---------------------------------------------------------------------------

def save_docx(text: str, path: Path) -> None:
    if Document is None:
        raise RuntimeError(
            "python-docx is required. Install it with: pip install python-docx"
        )

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    lines = [line.rstrip() for line in text.splitlines()]

    for index, line in enumerate(lines):
        stripped = line.strip()

        if not stripped:
            continue

        paragraph = doc.add_paragraph()

        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped)
            run.bold = True
            run.font.size = Pt(18)
            continue

        if stripped.upper() in SECTION_NAMES:
            run = paragraph.add_run(stripped.upper())
            run.bold = True
            run.font.size = Pt(11)
            continue

        if stripped.startswith(("-", "•", "*")):
            bullet = re.sub(r"^[\-•*]\s*", "", stripped)
            paragraph.style = "List Bullet"
            run = paragraph.add_run(bullet)
            run.font.size = Pt(10.5)
            continue

        run = paragraph.add_run(stripped)
        run.font.size = Pt(10.5)

    doc.save(path)


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    if not soffice:
        print(
            "[WARN] LibreOffice not found. "
            "DOCX generated, PDF skipped."
        )
        return None

    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        print(
            "[WARN] PDF conversion failed: "
            + (result.stderr.strip() or "unknown error")
        )
        return None

    pdf_path = output_dir / f"{docx_path.stem}.pdf"

    if pdf_path.exists():
        return pdf_path

    return None


# ---------------------------------------------------------------------------
# Main processing
# ---------------------------------------------------------------------------

def load_job_description(args: argparse.Namespace) -> str:
    if args.jd_file:
        path = Path(args.jd_file)

        if not path.exists():
            raise FileNotFoundError(
                f"Job description file not found: {path}"
            )

        return path.read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()

    if args.jd:
        return args.jd.strip()

    print()
    print("Paste the job description.")
    print("Type END on a separate line when finished.")
    print()

    lines = []

    while True:
        try:
            line = input()
        except EOFError:
            break

        if line.strip().upper() == "END":
            break

        lines.append(line)

    return "\n".join(lines).strip()


def run_pipeline(args: argparse.Namespace) -> None:

    resume_path = Path(args.resume)
    output_dir = Path(args.output_dir)

    if not resume_path.exists():
        raise FileNotFoundError(
            f"Resume file not found: {resume_path}"
        )

    output_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    print("=" * 78)
    print("HPPCS[01] - AI Resume Tailoring & ATS Optimization")
    print("=" * 78)
    print(f"Resume          : {resume_path}")
    print(f"Output          : {output_dir.resolve()}")
    print(f"Gemma model     : {args.gemma_model}")
    print(f"Review model    : {args.review_model}")
    print(f"Ollama          : {args.ollama_host}")
    print(f"Timeout         : {args.timeout}s")
    print()

    # 1. Resume extraction
    print("[1/7] Extracting resume text...")
    source_text = extract_resume_text(resume_path)

    (output_dir / "resume_source_text.txt").write_text(
        source_text,
        encoding="utf-8",
    )

    print(
        f"      Extracted {len(source_text):,} characters."
    )

    # 2. Resume structured extraction
    print("[2/7] Extracting structured resume facts with Gemma...")
    profile = extract_resume_profile(
        source_text,
        args.gemma_model,
        args.ollama_host,
        args.timeout,
    )

    (output_dir / "resume_profile.json").write_text(
        json.dumps(
            profile,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    # 3. JD parsing
    print("[3/7] Parsing job description with Gemma...")
    jd_text = load_job_description(args)

    if not jd_text:
        raise RuntimeError(
            "Job description is empty."
        )

    jd = parse_job_description(
        jd_text,
        args.gemma_model,
        args.ollama_host,
        args.timeout,
    )

    (output_dir / "job_description.json").write_text(
        json.dumps(
            jd,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    # 4. ATS matching
    print("[4/7] Calculating ATS keyword alignment...")
    ats = calculate_ats_analysis(
        profile,
        jd,
    )

    (output_dir / "ats_analysis.json").write_text(
        json.dumps(
            ats,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    print(
        f"      Keyword match : {ats['keyword_score']}%"
    )
    print(
        f"      Required match: {ats['required_skill_score']}%"
    )

    if ats["matched_keywords"]:
        print(
            "      Matched       : "
            + ", ".join(ats["matched_keywords"][:15])
        )

    if ats["missing_keywords"]:
        print(
            "      Missing       : "
            + ", ".join(ats["missing_keywords"][:15])
        )

    # 5. Tailored generation
    print("[5/7] Generating tailored ATS-friendly resume...")
    resume = generate_tailored_resume(
        profile,
        jd,
        ats,
        args.gemma_model,
        args.ollama_host,
        args.timeout,
    )

    # Optional user revision
    if args.revision:
        print("[INFO] Applying user revision request...")
        resume = revise_resume(
            resume,
            profile,
            jd,
            args.revision,
            args.gemma_model,
            args.ollama_host,
            args.timeout,
        )

    (output_dir / "tailored_resume.txt").write_text(
        resume,
        encoding="utf-8",
    )

    # 6. Review
    print("[6/7] Reviewing with Llama 3.2...")
    review = review_resume(
        resume,
        source_text,
        profile,
        jd,
        ats,
        args.review_model,
        args.ollama_host,
        args.timeout,
    )

    deterministic = validate_grounding(
        resume,
        profile,
    )

    review["python_validation"] = deterministic

    if deterministic["issues"]:
        review["recommendation"] = "REGENERATE"

    (output_dir / "final_review.json").write_text(
        json.dumps(
            review,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    # One regeneration only.
    if (
        str(review.get("recommendation", "")).upper()
        == "REGENERATE"
    ):
        print(
            "[WARN] Review found issues. "
            "Regenerating once with review feedback..."
        )

        issues = (
            review.get("issues", [])
            + deterministic.get("issues", [])
        )

        try:
            regenerated = generate_tailored_resume(
                profile,
                jd,
                ats,
                args.gemma_model,
                args.ollama_host,
                args.timeout,
                revision_instruction=(
                    "Correct these review issues without inventing facts:\n"
                    + "\n".join(
                        f"- {issue}"
                        for issue in issues[:12]
                    )
                ),
            )

            regenerated_validation = validate_grounding(
                regenerated,
                profile,
            )

            if (
                regenerated_validation["score"]
                >= deterministic["score"]
            ):
                resume = regenerated
                deterministic = regenerated_validation

                review = review_resume(
                    resume,
                    source_text,
                    profile,
                    jd,
                    ats,
                    args.review_model,
                    args.ollama_host,
                    args.timeout,
                )

                review["python_validation"] = deterministic

                (output_dir / "tailored_resume.txt").write_text(
                    resume,
                    encoding="utf-8",
                )

                (output_dir / "final_review.json").write_text(
                    json.dumps(
                        review,
                        indent=2,
                        ensure_ascii=False,
                    ),
                    encoding="utf-8",
                )

        except Exception as exc:
            print(
                f"[WARN] Regeneration failed: {exc}"
            )

    # 7. Output generation
    print("[7/7] Creating DOCX and PDF...")
    docx_path = (
        output_dir / "tailored_resume.docx"
    )

    save_docx(
        resume,
        docx_path,
    )

    pdf_path = convert_docx_to_pdf(
        docx_path,
        output_dir,
    )

    overall_score = review.get(
        "overall_score",
        ats["keyword_score"],
    )

    print()
    print("=" * 78)
    print("COMPLETED")
    print("=" * 78)
    print(
        f"ATS keyword match : {ats['keyword_score']}%"
    )
    print(
        f"Required match    : {ats['required_skill_score']}%"
    )
    print(
        f"Review score      : {overall_score}"
    )
    print(
        f"Review result     : "
        f"{review.get('recommendation', 'REVIEW')}"
    )
    print()
    print(f"DOCX              : {docx_path.resolve()}")

    if pdf_path:
        print(
            f"PDF               : {pdf_path.resolve()}"
        )
    else:
        print(
            "PDF               : Not generated "
            "(install LibreOffice)"
        )

    print()
    print("Generated files:")
    for path in sorted(output_dir.iterdir()):
        if path.is_file():
            print(f"  - {path.name}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Generate a truthful, ATS-tailored resume "
            "from PDF/DOC/DOCX + job description."
        )
    )

    parser.add_argument(
        "--resume",
        required=True,
        help="Path to candidate resume PDF, DOCX or DOC.",
    )

    parser.add_argument(
        "--jd",
        default="",
        help="Job description supplied directly as text.",
    )

    parser.add_argument(
        "--jd-file",
        default="",
        help="Path to a text file containing the job description.",
    )

    parser.add_argument(
        "--output-dir",
        default="./output",
        help="Directory for generated files.",
    )

    parser.add_argument(
        "--gemma-model",
        default=DEFAULT_GEMMA_MODEL,
        help="Ollama generation/extraction model.",
    )

    parser.add_argument(
        "--review-model",
        default=DEFAULT_REVIEW_MODEL,
        help="Ollama independent review model.",
    )

    parser.add_argument(
        "--ollama-host",
        default=DEFAULT_OLLAMA_HOST,
        help="Ollama API URL.",
    )

    parser.add_argument(
        "--timeout",
        type=int,
        default=300,
        help="HTTP timeout per Ollama request.",
    )

    parser.add_argument(
        "--revision",
        default="",
        help=(
            "Optional user revision request, e.g. "
            "'Make the summary more focused on cloud automation'."
        ),
    )

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    try:
        run_pipeline(args)
    except KeyboardInterrupt:
        print("\n[INFO] Cancelled by user.")
    except Exception as exc:
        print(f"\n[ERROR] {exc}")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
