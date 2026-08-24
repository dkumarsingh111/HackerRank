"""HPPCS01 - CV Creation using local open-source LLMs.

This pipeline processes unstructured candidate profiles in batch and creates
ATS-friendly CVs. It uses Gemma via Ollama for extraction/generation and
Llama 3.2 via Ollama for independent review.

Optional job-description support is included. When a JD is supplied, the CVs
are lightly tailored and ATS keyword alignment is reported per candidate.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List

try:
    import requests
except ImportError:
    requests = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError:
    Document = None
    WD_ALIGN_PARAGRAPH = None
    Inches = Pt = None


DEFAULT_GEMMA_MODEL = "gemma3:1b"
DEFAULT_LLAMA_MODEL = "llama3.2:1b"
DEFAULT_OLLAMA_HOST = "http://localhost:11434"

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".doc", ".rtf"}

SECTION_NAMES = {
    "SUMMARY": "PROFESSIONAL SUMMARY",
    "PROFESSIONAL SUMMARY": "PROFESSIONAL SUMMARY",
    "SKILLS": "SKILLS",
    "EXPERIENCE": "PROFESSIONAL EXPERIENCE",
    "WORK EXPERIENCE": "PROFESSIONAL EXPERIENCE",
    "PROFESSIONAL EXPERIENCE": "PROFESSIONAL EXPERIENCE",
    "EDUCATION": "EDUCATION",
    "ACHIEVEMENTS": "ACHIEVEMENTS",
    "PROJECTS": "PROJECTS",
    "CERTIFICATIONS": "CERTIFICATIONS",
    "INTERESTS": "INTERESTS",
}


def extract_pdf_text(path: Path) -> str:
    pages: List[str] = []

    if pdfplumber is not None:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(text.strip())

    elif PdfReader is not None:
        reader = PdfReader(str(path))
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text.strip())

    else:
        raise RuntimeError(
            "PDF support requires pdfplumber or pypdf. "
            "Install dependencies from requirements.txt."
        )

    if not pages:
        raise RuntimeError(
            f"No text could be extracted from {path.name}. "
            "The file may be scanned/image-only."
        )

    return "\n\n".join(pages)


def extract_docx_text(path: Path) -> str:
    if Document is None:
        raise RuntimeError(
            "python-docx is required for DOCX input. "
            "Install dependencies from requirements.txt."
        )

    doc = Document(path)
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    if not parts:
        raise RuntimeError(f"No text could be extracted from {path.name}.")

    return "\n".join(parts)


def extract_doc_text(path: Path) -> str:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError(
            "Legacy .doc input requires LibreOffice/soffice in PATH."
        )

    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                "txt:Text",
                "--outdir",
                tmp,
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )

        if result.returncode != 0:
            raise RuntimeError(
                result.stderr.strip() or "Could not convert DOC to text."
            )

        candidates = list(Path(tmp).glob("*.txt"))
        if not candidates:
            raise RuntimeError("LibreOffice produced no TXT output.")

        return candidates[0].read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()


def extract_rtf_text(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = re.sub(r"[{}]", " ", raw)
    return re.sub(r"\s+", " ", raw).strip()


def extract_text_from_file(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return path.read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    if suffix == ".doc":
        return extract_doc_text(path)

    if suffix == ".rtf":
        return extract_rtf_text(path)

    raise ValueError(
        f"Unsupported input format: {suffix}. "
        f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def discover_input_files(input_dir: Path, limit: int) -> List[Path]:
    candidates = [
        path
        for path in input_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and path.stem.lower().startswith("profile_")
    ]

    priority = {
        ".pdf": 0,
        ".docx": 1,
        ".doc": 2,
        ".rtf": 3,
        ".txt": 4,
        ".md": 5,
    }

    selected: Dict[str, Path] = {}
    for path in candidates:
        key = path.stem.lower()
        current = selected.get(key)
        if current is None or priority[path.suffix.lower()] < priority[current.suffix.lower()]:
            selected[key] = path

    return sorted(selected.values(), key=lambda item: item.name.lower())[:limit]


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def clean_list(value: Any) -> List[str]:
    if value is None:
        return []

    if isinstance(value, str):
        value = [value]

    if not isinstance(value, list):
        return []

    result: List[str] = []
    for item in value:
        if isinstance(item, dict):
            item = " | ".join(
                f"{key}: {val}"
                for key, val in item.items()
                if val not in (None, "")
            )

        item = clean_text(item)
        if item and item.casefold() not in {"none", "n/a", "not provided", "unknown"}:
            result.append(item)

    seen = set()
    unique: List[str] = []
    for item in result:
        key = item.casefold()
        if key not in seen:
            seen.add(key)
            unique.append(item)

    return unique


def extract_json(text: str) -> Dict[str, Any]:
    text = text.strip()
    text = re.sub(r"^```json\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.S)
        if not match:
            raise ValueError("LLM did not return valid JSON.")
        data = json.loads(match.group(0))

    if not isinstance(data, dict):
        raise ValueError("Expected a JSON object from the LLM.")

    return data


def call_ollama(
    model: str,
    prompt: str,
    host: str,
    timeout: int,
    num_predict: int,
    json_mode: bool = False,
) -> str:
    if requests is None:
        raise RuntimeError(
            "The 'requests' package is required. Install dependencies from requirements.txt."
        )

    payload: Dict[str, Any] = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "keep_alive": "5m",
        "options": {
            "temperature": 0.1,
            "top_p": 0.9,
            "num_predict": num_predict,
        },
    }

    if json_mode:
        payload["format"] = "json"

    response = requests.post(
        f"{host.rstrip('/')}/api/generate",
        json=payload,
        timeout=timeout,
    )
    response.raise_for_status()

    result = response.json().get("response", "")
    if not result.strip():
        raise RuntimeError(f"{model} returned an empty response.")

    return result.strip()


def extract_source_anchors(text: str) -> Dict[str, str]:
    email_match = re.search(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", text)
    email = email_match.group(0) if email_match else ""

    phone = ""
    phone_patterns = [
        r"(?<!\d)(\+\d{1,3}[\s().-]?\d{3,5}[\s().-]?\d{4,6})(?!\d)",
        r"(?<!\d)(\d{10})(?!\d)",
        r"(?<!\d)(\d{3}[\s-]\d{4})(?!\d)",
    ]
    for pattern in phone_patterns:
        match = re.search(pattern, text)
        if match:
            phone = match.group(1)
            break

    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]

    name = ""
    for line in lines[:8]:
        match = re.match(r"^(?:name|candidate)\s*[:\-]\s*(.+)$", line, flags=re.I)
        if match:
            name = clean_text(match.group(1))
            break

    if not name:
        for line in lines[:5]:
            if (
                1 <= len(line.split()) <= 5
                and "@" not in line
                and not re.search(r"\d", line)
                and not re.search(
                    r"\b(summary|profile|skills|experience|education)\b",
                    line,
                    flags=re.I,
                )
            ):
                name = line
                break

    location = ""
    for line in lines[:10]:
        match = re.search(r"(?:contact|location|city)\s*[:\-]\s*(.+)$", line, flags=re.I)
        if match:
            pieces = [clean_text(piece) for piece in re.split(r"[|,]", match.group(1))]
            for piece in pieces:
                if piece and "@" not in piece and not re.search(r"\+?\d[\d\s().-]{6,}", piece):
                    location = piece
            if location:
                break

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "location": location,
    }


def normalize_profile(data: Dict[str, Any], anchors: Dict[str, str]) -> Dict[str, Any]:
    contact = data.get("contact") or {}
    if not isinstance(contact, dict):
        contact = {}

    profile: Dict[str, Any] = {
        "name": clean_text(data.get("name")),
        "contact": {
            "email": clean_text(contact.get("email")),
            "phone": clean_text(contact.get("phone")),
            "location": clean_text(contact.get("location")),
        },
        "summary": clean_text(data.get("summary")),
        "education": clean_list(data.get("education")),
        "experience": clean_list(data.get("experience")),
        "skills": clean_list(data.get("skills")),
        "achievements": clean_list(data.get("achievements")),
        "projects": clean_list(data.get("projects")),
        "certifications": clean_list(data.get("certifications")),
        "interests": clean_list(data.get("interests")),
    }

    for field in ("name", "email", "phone", "location"):
        if anchors.get(field):
            if field == "name":
                profile["name"] = anchors[field]
            else:
                profile["contact"][field] = anchors[field]

    location = profile["contact"]["location"]
    if re.search(r"\b(school|high school|college|university|corp|inc|ltd|diploma|b\.tech)\b", location, re.I):
        profile["contact"]["location"] = ""

    return profile


def fallback_profile(text: str) -> Dict[str, Any]:
    anchors = extract_source_anchors(text)
    lines = [clean_text(line) for line in text.splitlines() if clean_text(line)]

    skills: List[str] = []
    education: List[str] = []
    achievements: List[str] = []
    projects: List[str] = []
    experience: List[str] = []
    summary = ""

    for line in lines:
        lower = line.lower()
        if not summary and re.match(r"(?i)^(summary|profile)\s*:", line):
            summary = clean_text(line.split(":", 1)[1])
        elif lower.startswith("skills:"):
            skills.extend([clean_text(item) for item in line.split(":", 1)[1].split(",")])
        elif lower.startswith("education:"):
            education.append(clean_text(line.split(":", 1)[1]))
        elif lower.startswith("achievements:"):
            achievements.append(clean_text(line.split(":", 1)[1]))
        elif lower.startswith("project:") or lower.startswith("projects:"):
            projects.append(clean_text(line.split(":", 1)[1]))
        elif " at " in lower or "worked" in lower:
            experience.append(line)

    return normalize_profile(
        {
            "name": anchors["name"],
            "contact": {
                "email": anchors["email"],
                "phone": anchors["phone"],
                "location": anchors["location"],
            },
            "summary": summary,
            "education": education,
            "experience": experience,
            "skills": skills,
            "achievements": achievements,
            "projects": projects,
        },
        anchors,
    )


def extract_profile(
    source_text: str,
    model: str,
    host: str,
    timeout: int,
) -> Dict[str, Any]:
    anchors = extract_source_anchors(source_text)

    prompt = f"""
You are a resume FACT EXTRACTION engine.

The source below may be an unstructured profile, short biography, paragraph,
plain text resume, or a partially formatted CV.

Extract ONLY facts explicitly supported by the SOURCE.

Return JSON only with exactly these keys:

{{
  "name": "",
  "contact": {{
    "email": "",
    "phone": "",
    "location": ""
  }},
  "summary": "",
  "education": [],
  "experience": [],
  "skills": [],
  "achievements": [],
  "projects": [],
  "certifications": [],
  "interests": []
}}

Rules:
- Never invent information.
- Do not infer a job title from an employer name.
- Do not infer responsibilities from a job title.
- Preserve employers and dates when stated.
- Preserve explicit achievements and metrics.
- Preserve explicit project names/details.
- Do not move responsibilities into achievements.
- If achievements, projects, certifications, or interests are not present, return [].
- Do not infer location.
- Do not add technologies not in the source.

SOURCE:
{source_text}
"""

    try:
        raw = call_ollama(
            model=model,
            prompt=prompt,
            host=host,
            timeout=timeout,
            num_predict=550,
            json_mode=True,
        )
        return normalize_profile(extract_json(raw), anchors)
    except Exception as exc:
        print(f"[WARN] Structured extraction fallback used: {exc}")
        return fallback_profile(source_text)


def parse_job_description(
    jd_text: str,
    model: str,
    host: str,
    timeout: int,
) -> Dict[str, Any]:
    prompt = f"""
You are a JOB DESCRIPTION extraction engine.

Extract the job posting into structured JSON.

Return exactly:

{{
  "job_title": "",
  "company": "",
  "requirements": [],
  "responsibilities": [],
  "preferred_qualifications": [],
  "keywords": []
}}

Rules:
- Extract only information present in the job description.
- Do not invent requirements.
- Keywords should include explicit technical skills, tools, certifications,
  job titles, domain terms and important phrases.
- Remove duplicate keywords.
- Keep requirements and responsibilities separate.

JOB DESCRIPTION:
{jd_text}
"""

    raw = call_ollama(
        model=model,
        prompt=prompt,
        host=host,
        timeout=timeout,
        num_predict=500,
        json_mode=True,
    )

    data = extract_json(raw)
    return {
        "job_title": clean_text(data.get("job_title")),
        "company": clean_text(data.get("company")),
        "requirements": clean_list(data.get("requirements")),
        "responsibilities": clean_list(data.get("responsibilities")),
        "preferred_qualifications": clean_list(data.get("preferred_qualifications")),
        "keywords": clean_list(data.get("keywords")),
    }


def build_profile_search_text(profile: Dict[str, Any]) -> str:
    values: List[str] = [
        profile.get("name", ""),
        profile.get("summary", ""),
        profile.get("contact", {}).get("location", ""),
    ]

    for field in (
        "education",
        "experience",
        "skills",
        "achievements",
        "projects",
        "certifications",
        "interests",
    ):
        values.extend(profile.get(field, []))

    return " ".join(values).casefold()


def keyword_match(keywords: List[str], search_text: str) -> tuple[List[str], List[str]]:
    matched: List[str] = []
    missing: List[str] = []

    for keyword in keywords:
        item = clean_text(keyword)
        if not item:
            continue
        if item.casefold() in search_text:
            matched.append(item)
        else:
            missing.append(item)

    return matched, missing


def calculate_ats_analysis(profile: Dict[str, Any], jd: Dict[str, Any]) -> Dict[str, Any]:
    search_text = build_profile_search_text(profile)
    keywords = clean_list(jd.get("keywords"))

    matched, missing = keyword_match(keywords, search_text)
    keyword_score = round((len(matched) / len(keywords)) * 100) if keywords else 100

    requirement_tokens = clean_list(
        re.findall(
            r"\b[A-Za-z][A-Za-z0-9+#./-]{1,30}\b",
            " ".join(jd.get("requirements", [])),
        )
    )
    required_matched, required_missing = keyword_match(requirement_tokens, search_text)
    required_score = (
        round((len(required_matched) / len(requirement_tokens)) * 100)
        if requirement_tokens
        else 100
    )

    return {
        "keyword_score": keyword_score,
        "required_skill_score": required_score,
        "matched_keywords": matched,
        "missing_keywords": missing,
        "matched_requirements": required_matched,
        "missing_requirements": required_missing,
        "recommendations": [
            f"Consider highlighting '{item}' only if it is genuinely supported by the candidate profile."
            for item in missing[:10]
        ],
    }


def generate_resume(
    profile: Dict[str, Any],
    source_text: str,
    model: str,
    host: str,
    timeout: int,
    jd: Dict[str, Any] | None = None,
    ats: Dict[str, Any] | None = None,
    revision_instruction: str = "",
) -> str:
    jd = jd or {}
    ats = ats or {}

    prompt = f"""
You are a professional resume writer.

Create a polished ATS-friendly CV using the CANDIDATE FACTS below.

The source may be completely unstructured.

Your task:
1. Reorganize information professionally.
2. Improve grammar and wording.
3. Keep every fact grounded in the source.
4. If a JOB DESCRIPTION is supplied, align relevant experience with it using
   only source-supported facts.
5. Do not fabricate missing skills.
6. Do not fabricate achievements, metrics, employers, titles, dates, or projects.
7. Do not create sections with no source data.

Required resume structure:

Actual candidate name on the first line
Actual phone/email/location line on the second line

PROFESSIONAL SUMMARY
SKILLS
PROFESSIONAL EXPERIENCE
EDUCATION

Optional only if source data exists:
ACHIEVEMENTS
PROJECTS
CERTIFICATIONS
INTERESTS

For experience:
- Use company, job title, and dates when supplied.
- Convert source responsibilities into concise professional bullets.
- Do not add unsupported responsibilities.

For the summary:
- Write 2-3 sentences.
- Base every claim on candidate facts.
- Do not claim unsupported expertise.

Critical formatting rules:
- Do not output literal placeholders such as "CANDIDATE NAME" or "Phone | Email | Location".
- The first line must be the candidate's actual name.
- The second line must contain the actual available contact details from the facts.

Return plain text only.
Do not return Markdown code fences or explanations.

CANDIDATE FACTS:
{json.dumps(profile, indent=2, ensure_ascii=False)}

SOURCE PROFILE:
{source_text}

JOB DESCRIPTION:
{json.dumps(jd, indent=2, ensure_ascii=False)}

ATS ANALYSIS:
{json.dumps(ats, indent=2, ensure_ascii=False)}

REVISION REQUEST:
{revision_instruction or "None"}
"""

    return call_ollama(
        model=model,
        prompt=prompt,
        host=host,
        timeout=timeout,
        num_predict=950,
        json_mode=False,
    ).strip()


def validate_resume(resume: str, profile: Dict[str, Any]) -> Dict[str, Any]:
    lower_resume = resume.casefold()
    issues: List[str] = []

    contact = profile.get("contact", {})
    protected_values = [
        ("name", profile.get("name", "")),
        ("email", contact.get("email", "")),
        ("phone", contact.get("phone", "")),
    ]

    missing = [
        label
        for label, value in protected_values
        if value and value.casefold() not in lower_resume
    ]
    if missing:
        issues.append("Missing source contact information: " + ", ".join(missing))

    for skill in profile.get("skills", []):
        if skill.casefold() not in lower_resume:
            issues.append(f"Source skill missing: {skill}")

    for project in profile.get("projects", []):
        words = [
            word.casefold()
            for word in re.findall(r"[A-Za-z0-9+#.-]+", project)
            if len(word) > 3
        ]
        if words:
            present = sum(word in lower_resume for word in words)
            if present / len(words) < 0.5:
                issues.append(f"Project may be missing: {project}")

    for achievement in profile.get("achievements", []):
        for number in re.findall(r"\d+(?:\.\d+)?%?", achievement):
            if number not in resume:
                issues.append(f"Achievement metric missing: {number}")

    for heading, field in (
        ("ACHIEVEMENTS", "achievements"),
        ("PROJECTS", "projects"),
        ("CERTIFICATIONS", "certifications"),
        ("INTERESTS", "interests"),
    ):
        if not profile.get(field) and re.search(rf"\b{heading}\b", resume, re.I):
            issues.append(f"{heading} section exists although source data is empty.")

    placeholders = re.findall(
        r"\[[^\]]+\]|\b(?:N/A|TBD|TO BE PROVIDED)\b",
        resume,
        flags=re.I,
    )
    if placeholders:
        issues.append("Placeholder text found: " + ", ".join(placeholders[:5]))

    score = max(0, 100 - min(100, len(issues) * 10))
    return {
        "score": score,
        "issues": issues,
        "pass": not issues,
    }


def review_resume(
    resume: str,
    source_text: str,
    profile: Dict[str, Any],
    model: str,
    host: str,
    timeout: int,
    jd: Dict[str, Any] | None = None,
    ats: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    jd_present = bool(jd)
    jd_rule = (
        "3. JOB DESCRIPTION when present"
        if jd_present
        else "3. No job description was provided for this run, so do not penalize the resume for missing JD alignment or ATS-analysis content."
    )
    prompt = f"""
You are an independent resume quality reviewer.

Review the generated resume against:
1. ORIGINAL SOURCE
2. STRUCTURED CANDIDATE FACTS
{jd_rule}

Return JSON only:

{{
  "overall_score": 0,
  "ats_score": 0,
  "factual_consistency": 0,
  "keyword_alignment": 0,
  "format_quality": 0,
  "missing_source_items": [],
  "unsupported_claims": [],
  "issues": [],
  "recommendation": "PASS"
}}

Use "REGENERATE" when there is a meaningful factual or structural problem.

Important:
- The source is the ultimate truth.
- Professional rewriting is allowed.
- Invented experience, skills, metrics, achievements, or projects are not allowed.
- Missing source achievements/projects are problems when they exist in the source.
- Empty source achievements/projects are not problems.
- If no job description was provided, never mark that as an issue.
- A missing JD keyword is not a hallucination; report it as missing rather than fabricate it.

ORIGINAL SOURCE:
{source_text}

STRUCTURED CANDIDATE FACTS:
{json.dumps(profile, indent=2, ensure_ascii=False)}

JOB DESCRIPTION:
{json.dumps(jd or {}, indent=2, ensure_ascii=False)}

ATS ANALYSIS:
{json.dumps(ats or {}, indent=2, ensure_ascii=False)}

GENERATED RESUME:
{resume}
"""

    try:
        return extract_json(
            call_ollama(
                model=model,
                prompt=prompt,
                host=host,
                timeout=timeout,
                num_predict=550,
                json_mode=True,
            )
        )
    except Exception as exc:
        return {
            "overall_score": 0,
            "ats_score": (ats or {}).get("keyword_score", 0),
            "factual_consistency": 0,
            "keyword_alignment": (ats or {}).get("keyword_score", 0),
            "format_quality": 0,
            "missing_source_items": [],
            "unsupported_claims": [],
            "issues": [f"Llama review failed: {exc}"],
            "recommendation": "REGENERATE",
        }


def build_contact_line(profile: Dict[str, Any]) -> str:
    contact = profile.get("contact", {})
    return " | ".join(
        value
        for value in (
            contact.get("phone", ""),
            contact.get("email", ""),
            contact.get("location", ""),
        )
        if clean_text(value)
    )


def repair_resume_output(resume: str, profile: Dict[str, Any]) -> str:
    lines = [line.rstrip() for line in resume.splitlines()]
    filtered: List[str] = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            filtered.append("")
            continue
        if stripped.upper() == "CANDIDATE NAME":
            continue
        if stripped.casefold() == "phone | email | location":
            continue
        filtered.append(line)

    while filtered and not filtered[0].strip():
        filtered.pop(0)

    name = clean_text(profile.get("name"))
    contact_line = build_contact_line(profile)

    if filtered and name:
        first = filtered[0].strip().casefold()
        if first in {"candidate name", "name"} or name.casefold() not in first:
            filtered.insert(0, name)
        else:
            filtered[0] = name
    elif name:
        filtered.insert(0, name)

    if contact_line:
        if len(filtered) < 2:
            filtered.insert(1 if filtered else 0, contact_line)
        else:
            second = filtered[1].strip().casefold()
            if second in {"phone | email | location", "email | phone | location"}:
                filtered[1] = contact_line
            elif profile.get("contact", {}).get("email", "").casefold() not in second:
                filtered.insert(1, contact_line)

    return "\n".join(filtered).strip()


def normalize_review(review: Dict[str, Any], resume: str, jd_present: bool) -> Dict[str, Any]:
    def clamp_int(value: Any) -> int:
        try:
            return max(0, min(100, int(float(value))))
        except (TypeError, ValueError):
            return 0

    issues = clean_list(review.get("issues"))
    if not jd_present:
        issues = [
            issue
            for issue in issues
            if "job description" not in issue.casefold()
            and "ats analysis" not in issue.casefold()
        ]

    resume_upper = resume.upper()
    section_checks = {
        "professional summary": "PROFESSIONAL SUMMARY" in resume_upper,
        "skills section": "SKILLS" in resume_upper,
        "professional experience": "PROFESSIONAL EXPERIENCE" in resume_upper,
        "education section": "EDUCATION" in resume_upper,
    }
    issues = [
        issue
        for issue in issues
        if not any(
            key in issue.casefold() and present
            for key, present in section_checks.items()
        )
    ]

    normalized = {
        "overall_score": clamp_int(review.get("overall_score")),
        "ats_score": clamp_int(review.get("ats_score")),
        "factual_consistency": clamp_int(review.get("factual_consistency")),
        "keyword_alignment": clamp_int(review.get("keyword_alignment")),
        "format_quality": clamp_int(review.get("format_quality")),
        "missing_source_items": clean_list(review.get("missing_source_items")),
        "unsupported_claims": clean_list(review.get("unsupported_claims")),
        "issues": issues,
        "recommendation": clean_text(review.get("recommendation")).upper() or "REVIEW",
    }

    if normalized["recommendation"] not in {"PASS", "REGENERATE", "REVIEW"}:
        normalized["recommendation"] = "REVIEW"

    return normalized


def clean_resume_text(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:text|markdown)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text).strip()
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    return text


def save_docx(text: str, path: Path) -> None:
    if Document is None:
        raise RuntimeError(
            "python-docx is required. Install dependencies from requirements.txt."
        )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    lines = [line.rstrip() for line in text.splitlines()]

    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            continue

        paragraph = doc.add_paragraph()
        normalized_heading = stripped.upper().rstrip(":")

        if index == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(18)
            continue

        if normalized_heading in SECTION_NAMES:
            run = paragraph.add_run(SECTION_NAMES[normalized_heading])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(11)
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(3)
            continue

        if stripped.startswith(("-", "•", "*")):
            bullet = re.sub(r"^[\-•*]\s*", "", stripped)
            paragraph.style = "List Bullet"
            run = paragraph.add_run(bullet)
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            continue

        run = paragraph.add_run(stripped)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)

    doc.save(path)


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None

    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        return None

    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    return pdf_path if pdf_path.exists() else None


def load_job_description(args: argparse.Namespace) -> str:
    if args.jd_file:
        path = Path(args.jd_file)
        if not path.exists():
            raise FileNotFoundError(f"Job description file not found: {path}")
        return path.read_text(encoding="utf-8", errors="replace").strip()

    if args.jd:
        return args.jd.strip()

    return ""


def process_profile(
    file_path: Path,
    output_dir: Path,
    args: argparse.Namespace,
    jd_text: str = "",
    jd_data: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    print(f"\n[INFO] Processing {file_path.name}")

    source_text = extract_text_from_file(file_path)
    profile = extract_profile(
        source_text,
        args.gemma_model,
        args.ollama_host,
        args.timeout,
    )

    ats: Dict[str, Any] = {}
    if jd_text and jd_data is not None:
        ats = calculate_ats_analysis(profile, jd_data)

    resume = clean_resume_text(
        generate_resume(
            profile,
            source_text,
            args.gemma_model,
            args.ollama_host,
            args.timeout,
            jd=jd_data,
            ats=ats,
        )
    )
    resume = repair_resume_output(resume, profile)

    review = review_resume(
        resume,
        source_text,
        profile,
        args.llama_model,
        args.ollama_host,
        args.timeout,
        jd=jd_data,
        ats=ats,
    )
    review = normalize_review(review, resume, bool(jd_text))

    deterministic = validate_resume(resume, profile)
    review["python_validation"] = deterministic
    if deterministic["issues"]:
        review["recommendation"] = "REGENERATE"

    if str(review.get("recommendation", "")).upper() == "REGENERATE":
        print("[WARN] Review requested regeneration; retrying once...")
        issues = review.get("issues", []) + deterministic.get("issues", [])
        regenerated = clean_resume_text(
            generate_resume(
                profile,
                source_text,
                args.gemma_model,
                args.ollama_host,
                args.timeout,
                jd=jd_data,
                ats=ats,
                revision_instruction=(
                    "Correct these review issues without inventing facts:\n"
                    + "\n".join(f"- {issue}" for issue in issues[:12])
                ),
            )
        )
        regenerated = repair_resume_output(regenerated, profile)
        regenerated_validation = validate_resume(regenerated, profile)
        if regenerated_validation["score"] >= deterministic["score"]:
            resume = regenerated
            deterministic = regenerated_validation
            review = review_resume(
                resume,
                source_text,
                profile,
                args.llama_model,
                args.ollama_host,
                args.timeout,
                jd=jd_data,
                ats=ats,
            )
            review = normalize_review(review, resume, bool(jd_text))
            review["python_validation"] = deterministic

    if review.get("recommendation") == "REVIEW":
        review["recommendation"] = "PASS" if deterministic["pass"] else "REGENERATE"

    if (
        review.get("recommendation") == "REGENERATE"
        and not review.get("issues")
        and not review.get("unsupported_claims")
        and deterministic["pass"]
    ):
        review["recommendation"] = "PASS"

    stem = file_path.stem
    (output_dir / f"{stem}_source_text.txt").write_text(source_text, encoding="utf-8")
    (output_dir / f"{stem}_structured.json").write_text(
        json.dumps(profile, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    (output_dir / f"{stem}_final_cv.txt").write_text(resume, encoding="utf-8")
    (output_dir / f"{stem}_review.json").write_text(
        json.dumps(review, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    if ats:
        (output_dir / f"{stem}_ats_analysis.json").write_text(
            json.dumps(ats, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )

    docx_path = output_dir / f"{stem}_final_cv.docx"
    save_docx(resume, docx_path)
    pdf_path = convert_docx_to_pdf(docx_path, output_dir)

    return {
        "input": file_path.name,
        "output_cv": docx_path.name,
        "output_pdf": pdf_path.name if pdf_path else "",
        "structured_output": f"{stem}_structured.json",
        "review_output": f"{stem}_review.json",
        "ats_output": f"{stem}_ats_analysis.json" if ats else "",
        "keyword_score": ats.get("keyword_score", 0) if ats else 0,
        "required_skill_score": ats.get("required_skill_score", 0) if ats else 0,
        "review_score": review.get("overall_score", 0),
        "status": review.get("recommendation", "REVIEW"),
        "issues": review.get("issues", []),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Batch-generate ATS-friendly CVs from unstructured candidate profiles "
            "using local Ollama models."
        )
    )

    parser.add_argument(
        "--input-dir",
        default=".",
        help="Directory containing profile_*.txt/pdf/docx/doc/rtf files.",
    )
    parser.add_argument(
        "--output-dir",
        default=".",
        help="Directory for generated CVs and evaluation artifacts.",
    )
    parser.add_argument(
        "--gemma-model",
        default=DEFAULT_GEMMA_MODEL,
        help="Ollama model for extraction and CV generation.",
    )
    parser.add_argument(
        "--llama-model",
        default=DEFAULT_LLAMA_MODEL,
        help="Ollama model for independent CV review.",
    )
    parser.add_argument(
        "--review-model",
        default="",
        help="Backward-compatible alias for --llama-model.",
    )
    parser.add_argument(
        "--ollama-host",
        default=DEFAULT_OLLAMA_HOST,
        help="Base URL of the Ollama API.",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=240,
        help="HTTP timeout per Ollama request.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=10,
        help="Maximum number of profiles to process.",
    )
    parser.add_argument(
        "--jd",
        default="",
        help="Optional job description text used for tailoring.",
    )
    parser.add_argument(
        "--jd-file",
        default="",
        help="Optional path to a text file containing a job description.",
    )

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.review_model and not args.llama_model:
        args.llama_model = args.review_model
    elif args.review_model:
        args.llama_model = args.review_model

    input_dir = Path(args.input_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if not input_dir.exists():
        raise SystemExit(f"Input directory not found: {input_dir}")

    files = discover_input_files(input_dir, args.limit)
    if not files:
        raise SystemExit(
            "No supported profile files found. Expected files such as "
            "profile_01.txt, profile_02.pdf, or profile_03.docx."
        )

    jd_text = load_job_description(args)
    jd_data = None
    if jd_text:
        print("[INFO] Parsing job description with Gemma...")
        jd_data = parse_job_description(
            jd_text,
            args.gemma_model,
            args.ollama_host,
            args.timeout,
        )
        (output_dir / "job_description.json").write_text(
            json.dumps(jd_data, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )

    print("=" * 78)
    print("HPPCS[01] - CV Creation using LLMs")
    print("=" * 78)
    print(f"Input directory : {input_dir.resolve()}")
    print(f"Output directory: {output_dir.resolve()}")
    print(f"Profiles found  : {len(files)}")
    print(f"Gemma model     : {args.gemma_model}")
    print(f"Llama model     : {args.llama_model}")
    print(f"Ollama host     : {args.ollama_host}")
    print(f"Job description : {'provided' if jd_text else 'not provided'}")

    results: List[Dict[str, Any]] = []
    for file_path in files:
        try:
            results.append(
                process_profile(
                    file_path=file_path,
                    output_dir=output_dir,
                    args=args,
                    jd_text=jd_text,
                    jd_data=jd_data,
                )
            )
        except Exception as exc:
            print(f"[ERROR] Failed to process {file_path.name}: {exc}")
            results.append(
                {
                    "input": file_path.name,
                    "output_cv": "",
                    "output_pdf": "",
                    "structured_output": "",
                    "review_output": "",
                    "ats_output": "",
                    "keyword_score": 0,
                    "required_skill_score": 0,
                    "review_score": 0,
                    "status": "ERROR",
                    "issues": [str(exc)],
                }
            )

    summary = {
        "project": "HPPCS01 - CV Creation using LLMs",
        "models": {
            "generation": args.gemma_model,
            "review": args.llama_model,
        },
        "job_description_used": bool(jd_text),
        "profiles_processed": len(results),
        "successful_profiles": sum(1 for item in results if item["status"] != "ERROR"),
        "results": results,
    }

    if jd_text:
        keyword_scores = [item["keyword_score"] for item in results if item["status"] != "ERROR"]
        if keyword_scores:
            summary["average_keyword_score"] = round(sum(keyword_scores) / len(keyword_scores), 2)

    (output_dir / "evaluation_summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    print("\n" + "=" * 78)
    print("COMPLETED")
    print("=" * 78)
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
